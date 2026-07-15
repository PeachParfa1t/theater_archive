from io import BytesIO
from flask import request, send_file
from flask_login import login_required
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from docx import Document as DocxDocument
from app import Tour
from . import reports_bp, safe_filename, XLSX_MIME, DOCX_MIME

def _tour_artist_rows(t):
    """One row per (artist, production, role) the artist is credited with in a production
    that's part of this tour; falls back to a single placeholder row for artists linked to
    the tour with no matching cast entry in any of its productions."""
    prod_ids = {link.production_id for link in t.production_links}
    rows = []
    for link in t.artist_links:
        artist = link.artist
        matches = [ce for ce in artist.cast_entries if ce.production_id in prod_ids]
        if matches:
            for ce in matches:
                rows.append((artist, ce.production.name, ce.role_name or '—'))
        else:
            rows.append((artist, '—', '—'))
    return rows

def _tours_report_docx(title, headers, rows, tours):
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val) if val not in (None, '') else '—'

    for t in tours:
        doc.add_heading(f'{t.city}, {t.year}', level=2)

        doc.add_heading('Постановки', level=3)
        if t.production_links:
            pt = doc.add_table(rows=1, cols=4)
            pt.style = 'Light Grid Accent 1'
            hdr = pt.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = 'Название', 'Жанр', 'Сцена', 'Продолжительность'
            for link in t.production_links:
                p = link.production
                cells = pt.add_row().cells
                cells[0].text = p.name
                cells[1].text = p.genre
                cells[2].text = p.stage or '—'
                cells[3].text = p.duration_display if (p.duration_hours or p.duration_minutes) else '—'
        else:
            doc.add_paragraph('Нет данных.')

        doc.add_heading('Артисты', level=3)
        artist_rows = _tour_artist_rows(t)
        if artist_rows:
            at = doc.add_table(rows=1, cols=5)
            at.style = 'Light Grid Accent 1'
            hdr = at.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text, hdr[4].text = 'ФИО', 'Должность', 'Звание', 'Постановка', 'Роль'
            for artist, prod_name, role in artist_rows:
                cells = at.add_row().cells
                cells[0].text = artist.full_name
                cells[1].text = artist.position or '—'
                cells[2].text = artist.title or '—'
                cells[3].text = prod_name
                cells[4].text = role
        else:
            doc.add_paragraph('Нет данных.')

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'{safe_filename(title)}.docx', mimetype=DOCX_MIME)

def _tours_report_xlsx(headers, rows, tours, title):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Гастроли'
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(str(h)) + 4)

    ws2 = wb.create_sheet('Постановки по гастролям')
    ws2.append(['Гастроль', 'Название', 'Жанр', 'Сцена', 'Продолжительность'])
    for t in tours:
        label = f'{t.city}, {t.year}'
        for link in t.production_links:
            p = link.production
            ws2.append([label, p.name, p.genre, p.stage or '',
                       p.duration_display if (p.duration_hours or p.duration_minutes) else ''])
    for i in range(1, 6):
        ws2.column_dimensions[get_column_letter(i)].width = 26

    ws3 = wb.create_sheet('Артисты по гастролям')
    ws3.append(['Гастроль', 'ФИО', 'Должность', 'Звание', 'Постановка', 'Роль'])
    for t in tours:
        label = f'{t.city}, {t.year}'
        for artist, prod_name, role in _tour_artist_rows(t):
            ws3.append([label, artist.full_name, artist.position or '', artist.title or '', prod_name, role])
    for i in range(1, 7):
        ws3.column_dimensions[get_column_letter(i)].width = 26

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'{safe_filename(title)}.xlsx', mimetype=XLSX_MIME)

# ---------- Гастроли: фильтруемый список ----------
@reports_bp.route('/tours/export')
@login_required
def export_tours():
    fmt = request.args.get('format', 'xlsx')
    city = request.args.get('city', '').strip()
    year_from = request.args.get('year_from', '').strip()
    year_to = request.args.get('year_to', '').strip()
    tour_type = request.args.get('tour_type', '').strip()
    artist_q = request.args.get('artist', '').strip()
    production_q = request.args.get('production', '').strip()

    tours = Tour.query.order_by(Tour.year.desc(), Tour.city).all()

    def matches(t):
        if city and t.city != city:
            return False
        if year_from and t.year < int(year_from):
            return False
        if year_to and t.year > int(year_to):
            return False
        if tour_type and t.tour_type_name != tour_type:
            return False
        # Matched in Python, not via SQL LIKE/ilike: SQLite's lower() only folds ASCII
        # a-z and leaves Cyrillic untouched (see utils.get_or_create_libretto_role).
        if artist_q:
            target = artist_q.lower()
            if not any(target in link.artist.full_name.lower() for link in t.artist_links):
                return False
        if production_q:
            target = production_q.lower()
            if not any(target in link.production.name.lower() for link in t.production_links):
                return False
        return True

    filtered = [t for t in tours if matches(t)]

    headers = ['№', 'Город', 'Год', 'Номер поездки', 'Тип', 'Постановки', 'Количество артистов']
    rows = [[
        i + 1, t.city, t.year, t.tour_number or '—', t.tour_type_name,
        ', '.join(link.production.name for link in t.production_links) or '—',
        len(t.artist_links),
    ] for i, t in enumerate(filtered)]

    title = 'Отчёт по гастролям'
    if fmt == 'docx':
        return _tours_report_docx(title, headers, rows, filtered)
    return _tours_report_xlsx(headers, rows, filtered, title)
