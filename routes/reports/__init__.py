import re
from io import BytesIO
from flask import Blueprint, render_template, send_file
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from docx import Document as DocxDocument
from app import (Production, Tour, TourType, FestivalStatus, CompetitionStatus,
                 AwardLevel, report_generation_required)

reports_bp = Blueprint('reports', __name__, url_prefix='/reports')

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

def safe_filename(name):
    return re.sub(r'[\\/:*?"<>|]', '_', name).strip() or 'report'

def xlsx_response(headers, rows, filename):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(str(h)) + 4)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'{safe_filename(filename)}.xlsx', mimetype=XLSX_MIME)

def docx_table_response(title, headers, rows, filename):
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val) if val not in (None, '') else '—'
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'{safe_filename(filename)}.docx', mimetype=DOCX_MIME)

@reports_bp.route('/')
@report_generation_required
def index():
    genres = sorted({p.genre for p in Production.query.all() if p.genre})
    tour_cities = sorted({t.city for t in Tour.query.all() if t.city})
    tour_types = TourType.query.order_by(TourType.name).all()
    festival_statuses = FestivalStatus.query.order_by(FestivalStatus.name).all()
    competition_statuses = CompetitionStatus.query.order_by(CompetitionStatus.name).all()
    award_levels = AwardLevel.query.order_by(AwardLevel.name).all()
    return render_template('reports/index.html', genres=genres,
                           stages=Production.STAGES, age_ratings=Production.AGE_RATINGS,
                           tour_cities=tour_cities, tour_types=tour_types,
                           festival_statuses=festival_statuses,
                           competition_statuses=competition_statuses,
                           award_levels=award_levels)

# Submodules attach their routes to reports_bp on import; must come after the blueprint
# and shared helpers above are defined.
from . import artist_reports, production_reports, tour_reports, festival_reports, competition_reports
