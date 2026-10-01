from io import BytesIO
from flask import request, send_file
from flask_login import login_required
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from docx import Document as DocxDocument
from app import db, Production
from audit import commit_with_audit
from . import reports_bp, safe_filename, xlsx_response, docx_table_response, XLSX_MIME, DOCX_MIME

# ---------- Постановки: фильтруемый список ----------
@reports_bp.route('/productions/export')
@login_required
def export_productions():
    fmt = request.args.get('format', 'xlsx')
    status = request.args.get('status', '')
    genre = request.args.get('genre', '').strip()
    year_from = request.args.get('year_from', '').strip()
    year_to = request.args.get('year_to', '').strip()
    basis = request.args.get('basis', '').strip()
    stage = request.args.get('stage', '').strip()
    age_rating = request.args.get('age_rating', '').strip()
    removal_year_from = request.args.get('removal_year_from', '').strip()
    removal_year_to = request.args.get('removal_year_to', '').strip()

    query = Production.query
    if status:
        query = query.filter(Production.status == status)
    if genre:
        query = query.filter(Production.genre.ilike(f'%{genre}%'))
    if year_from:
        query = query.filter(Production.premiere_year >= int(year_from))
    if year_to:
        query = query.filter(Production.premiere_year <= int(year_to))
    if basis:
        query = query.filter(db.or_(
            Production.literary_basis.ilike(f'%{basis}%'),
            Production.literary_basis_author.ilike(f'%{basis}%'),
        ))
    if stage:
        query = query.filter(Production.stage == stage)
    if age_rating:
        query = query.filter(Production.age_rating == age_rating)
    if removal_year_from:
        query = query.filter(Production.removal_year >= int(removal_year_from))
    if removal_year_to:
        query = query.filter(Production.removal_year <= int(removal_year_to))
    productions = query.order_by(Production.premiere_year.desc()).all()

    headers = ['Название', 'Статус', 'Жанр', 'Премьера', 'Актов',
               'Авторы музыки', 'Авторы либретто', 'Литературная основа', 'Автор основы',
               'Сцена', 'Возрастной ценз', 'Продолжительность', 'Год снятия']
    rows = [[
        p.name, p.status_display, p.genre, p.premiere_display, p.acts_count or '',
        p.music_authors_display, p.libretto_authors_display,
        p.literary_basis or '', p.literary_basis_author or '',
        p.stage or '', p.age_rating or '',
        p.duration_display if (p.duration_hours or p.duration_minutes) else '',
        p.removal_year or '',
    ] for p in productions]

    title = 'Отчёт по постановкам'
    if fmt == 'docx':
        response = docx_table_response(title, headers, rows, title)
        actual_format = 'DOCX'
    else:
        response = xlsx_response(headers, rows, title)
        actual_format = 'XLSX'
    commit_with_audit('report', 'Отчёт по постановкам', title,
                      details=f'Формат: {actual_format}; строк: {len(rows)}')
    return response

# ---------- Полная информация по одной постановке ----------
REPORT_SECTIONS = {'cast', 'staging', 'documents', 'materials'}

def _production_docx(p, sections):
    doc = DocxDocument()
    doc.add_heading(p.name, level=1)
    doc.add_paragraph(f'Статус: {p.status_display}')
    doc.add_paragraph(f'Жанр: {p.genre}')
    doc.add_paragraph(f'Премьера: {p.premiere_display}')
    if p.acts_count:
        doc.add_paragraph(f'Количество актов: {p.acts_count}')
    if p.stage:
        doc.add_paragraph(f'Сцена: {p.stage}')
    if p.age_rating:
        doc.add_paragraph(f'Возрастной ценз: {p.age_rating}')
    if p.duration_hours or p.duration_minutes:
        doc.add_paragraph(f'Продолжительность: {p.duration_display}')
    if p.removal_year:
        doc.add_paragraph(f'Год снятия: {p.removal_year}')
    if p.music_authors:
        doc.add_paragraph(f'Авторы музыки: {p.music_authors_display}')
    if p.libretto_authors:
        doc.add_paragraph(f'Авторы либретто: {p.libretto_authors_display}')
    if p.literary_basis:
        line = f'Литературная основа: {p.literary_basis}'
        if p.literary_basis_author:
            line += f' (автор: {p.literary_basis_author})'
        doc.add_paragraph(line)

    if 'cast' in sections:
        doc.add_heading('Состав исполнителей', level=2)
        if p.cast_entries:
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Light Grid Accent 1'
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text = 'Артист', 'Роль', 'Годы участия'
            for ce in p.cast_entries:
                cells = table.add_row().cells
                cells[0].text = ce.artist.full_name
                cells[1].text = ce.role_name or '—'
                cells[2].text = ce.years_display or '—'
        else:
            doc.add_paragraph('Нет данных.')

    if 'staging' in sections:
        doc.add_heading('Постановочная группа', level=2)
        if p.staging_group:
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Light Grid Accent 1'
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text = 'ФИО', 'Должность', 'Годы жизни'
            for pd in p.staging_group:
                cells = table.add_row().cells
                cells[0].text = pd.director.full_name
                cells[1].text = pd.position_display
                cells[2].text = pd.director.life_years or '—'
        else:
            doc.add_paragraph('Нет данных.')

    if 'documents' in sections:
        doc.add_heading('Документы', level=2)
        if p.documents:
            for d in p.documents:
                doc.add_paragraph(f'{d.doc_type_display}: {d.title or d.file_name}', style='List Bullet')
        else:
            doc.add_paragraph('Нет данных.')

    if 'materials' in sections:
        doc.add_heading('Материалы', level=2)
        if p.materials:
            for m in p.materials:
                doc.add_paragraph(f'{m.type_display}: {m.title or m.file_name or m.url}', style='List Bullet')
        else:
            doc.add_paragraph('Нет данных.')

    return doc

def _production_xlsx(p, sections):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Информация'
    ws.append(['Поле', 'Значение'])
    ws.append(['Название', p.name])
    ws.append(['Статус', p.status_display])
    ws.append(['Жанр', p.genre])
    ws.append(['Премьера', p.premiere_display])
    ws.append(['Актов', p.acts_count or ''])
    ws.append(['Авторы музыки', p.music_authors_display])
    ws.append(['Авторы либретто', p.libretto_authors_display])
    ws.append(['Литературная основа', p.literary_basis or ''])
    ws.append(['Автор основы', p.literary_basis_author or ''])
    ws.append(['Сцена', p.stage or ''])
    ws.append(['Возрастной ценз', p.age_rating or ''])
    ws.append(['Продолжительность', p.duration_display if (p.duration_hours or p.duration_minutes) else ''])
    ws.append(['Год снятия', p.removal_year or ''])
    for col, width in (('A', 22), ('B', 50)):
        ws.column_dimensions[col].width = width

    sheets = []

    if 'cast' in sections:
        ws2 = wb.create_sheet('Состав')
        ws2.append(['Артист', 'Роль', 'Годы участия'])
        for ce in p.cast_entries:
            ws2.append([ce.artist.full_name, ce.role_name or '', ce.years_display or ''])
        sheets.append(ws2)

    if 'staging' in sections:
        ws3 = wb.create_sheet('Постановочная группа')
        ws3.append(['ФИО', 'Должность', 'Годы жизни'])
        for pd in p.staging_group:
            ws3.append([pd.director.full_name, pd.position_display, pd.director.life_years or ''])
        sheets.append(ws3)

    if 'documents' in sections:
        ws4 = wb.create_sheet('Документы')
        ws4.append(['Тип', 'Название', 'Файл'])
        for d in p.documents:
            ws4.append([d.doc_type_display, d.title or '', d.file_name])
        sheets.append(ws4)

    if 'materials' in sections:
        ws5 = wb.create_sheet('Материалы')
        ws5.append(['Тип', 'Название', 'Файл/ссылка'])
        for m in p.materials:
            ws5.append([m.type_display, m.title or '', m.file_name or m.url or ''])
        sheets.append(ws5)

    for sheet in sheets:
        for i, _ in enumerate(sheet[1], 1):
            sheet.column_dimensions[get_column_letter(i)].width = 28

    return wb

@reports_bp.route('/production/<int:pid>/export')
@login_required
def export_production(pid):
    p = db.get_or_404(Production, pid)
    fmt = request.args.get('format', 'docx')
    sections_param = request.args.get('sections', '').strip()
    sections = {s for s in sections_param.split(',') if s in REPORT_SECTIONS} if sections_param else set(REPORT_SECTIONS)
    filename = safe_filename(p.name)
    if fmt == 'xlsx':
        wb = _production_xlsx(p, sections)
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        response = send_file(buf, as_attachment=True, download_name=f'{filename}.xlsx', mimetype=XLSX_MIME)
        actual_format = 'XLSX'
    else:
        doc = _production_docx(p, sections)
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        response = send_file(buf, as_attachment=True, download_name=f'{filename}.docx', mimetype=DOCX_MIME)
        actual_format = 'DOCX'
    commit_with_audit('report', 'Отчёт по постановке', p.name, p.id,
                      f'Формат: {actual_format}; разделы: {", ".join(sorted(sections))}')
    return response
