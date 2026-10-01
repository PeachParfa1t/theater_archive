from io import BytesIO
from flask import request, send_file
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from docx import Document as DocxDocument
from app import CastEntry, Artist, Production, report_download_required
from audit import commit_with_audit
from . import reports_bp, safe_filename, XLSX_MIME, DOCX_MIME

# ---------- Артисты: участие в постановках за период ----------
ARTIST_REPORT_SECTIONS = {'tours', 'competitions'}

def _artist_report_docx(title, headers, rows, artists, sections):
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val) if val not in (None, '') else '—'

    if 'tours' in sections:
        for artist in artists:
            doc.add_heading(f'Гастроли — {artist.full_name}', level=2)
            if artist.tour_links:
                t2 = doc.add_table(rows=1, cols=3)
                t2.style = 'Light Grid Accent 1'
                hdr = t2.rows[0].cells
                hdr[0].text, hdr[1].text, hdr[2].text = 'Город', 'Год', 'Постановки'
                for link in artist.tour_links:
                    cells = t2.add_row().cells
                    cells[0].text = link.tour.city
                    cells[1].text = str(link.tour.year)
                    cells[2].text = ', '.join(l.production.name for l in link.tour.production_links) or '—'
            else:
                doc.add_paragraph('Нет данных.')

    if 'competitions' in sections:
        for artist in artists:
            doc.add_heading(f'Конкурсы — {artist.full_name}', level=2)
            if artist.competition_links:
                t3 = doc.add_table(rows=1, cols=3)
                t3.style = 'Light Grid Accent 1'
                hdr = t3.rows[0].cells
                hdr[0].text, hdr[1].text, hdr[2].text = 'Название', 'Год', 'Уровень награды'
                for link in artist.competition_links:
                    cells = t3.add_row().cells
                    cells[0].text = link.competition.name
                    cells[1].text = str(link.competition.year)
                    cells[2].text = link.award_level_name
            else:
                doc.add_paragraph('Нет данных.')

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'{safe_filename(title)}.docx', mimetype=DOCX_MIME)

def _artist_report_xlsx(headers, rows, artists, sections, title):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Репертуар'
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(str(h)) + 4)

    if 'tours' in sections:
        ws2 = wb.create_sheet('Гастроли')
        ws2.append(['Артист', 'Город', 'Год', 'Постановки'])
        for artist in artists:
            for link in artist.tour_links:
                ws2.append([artist.full_name, link.tour.city, link.tour.year,
                           ', '.join(l.production.name for l in link.tour.production_links) or ''])
        for i in range(1, 5):
            ws2.column_dimensions[get_column_letter(i)].width = 28

    if 'competitions' in sections:
        ws3 = wb.create_sheet('Конкурсы')
        ws3.append(['Артист', 'Название', 'Год', 'Уровень награды'])
        for artist in artists:
            for link in artist.competition_links:
                ws3.append([artist.full_name, link.competition.name, link.competition.year, link.award_level_name])
        for i in range(1, 5):
            ws3.column_dimensions[get_column_letter(i)].width = 28

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'{safe_filename(title)}.xlsx', mimetype=XLSX_MIME)

@reports_bp.route('/artists/export')
@report_download_required
def export_artists():
    fmt = request.args.get('format', 'xlsx')
    q = request.args.get('q', '').strip()
    year_from = request.args.get('year_from', '').strip()
    year_to = request.args.get('year_to', '').strip()
    sections = {s for s in request.args.getlist('sections') if s in ARTIST_REPORT_SECTIONS}

    query = CastEntry.query.join(Artist).join(Production)
    if q:
        query = query.filter(Artist.full_name.ilike(f'%{q}%'))
    entries = query.order_by(Artist.full_name, Production.premiere_year).all()

    yf = int(year_from) if year_from else None
    yt = int(year_to) if year_to else None

    def overlaps(ce):
        if not yf and not yt:
            return True
        ce_from = ce.year_from or ce.production.premiere_year
        ce_to = ce.year_to or ce_from
        if yf and ce_to < yf:
            return False
        if yt and ce_from > yt:
            return False
        return True

    def stage_director_name(production):
        """The production's Режиссёр-постановщик, if one is assigned in the staging group."""
        for pd in production.staging_group:
            codes = [pp.position for pp in pd.positions] or pd.director.position_codes
            if 'director' in codes:
                return pd.director.full_name
        return ''

    filtered = [ce for ce in entries if overlaps(ce)]

    headers = ['№', 'Название спектакля', 'Автор (композитор)', 'Роль', 'Год выпуска', 'Режиссёр-постановщик']
    rows = [[
        i + 1, ce.production.name, ce.production.music_authors_display, ce.role_name or '—',
        ce.production.premiere_year, stage_director_name(ce.production) or '—',
    ] for i, ce in enumerate(filtered)]

    title = 'Репертуарный лист артиста'
    if q:
        title += f' — {q}'

    artists_in_report = sorted({ce.artist for ce in filtered}, key=lambda a: a.full_name)

    if fmt == 'docx':
        response = _artist_report_docx(title, headers, rows, artists_in_report, sections)
        actual_format = 'DOCX'
    else:
        response = _artist_report_xlsx(headers, rows, artists_in_report, sections, title)
        actual_format = 'XLSX'
    commit_with_audit('report', 'Отчёт по артистам', title,
                      details=f'Формат: {actual_format}; строк: {len(rows)}')
    return response
