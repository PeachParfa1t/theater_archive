from io import BytesIO
from flask import request, send_file
from flask_login import login_required
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from docx import Document as DocxDocument
from docx.shared import Pt
from app import FestivalEdition
from . import reports_bp, safe_filename, XLSX_MIME, DOCX_MIME

def _festivals_report_docx(title, headers, rows, editions):
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val) if val not in (None, '') else '—'

    for e in editions:
        doc.add_heading(f'{e.festival.name}, {e.year}', level=2)

        doc.add_heading('Описание', level=3)
        if e.description:
            p = doc.add_paragraph()
            lines = e.description.split('\n')
            for i, line in enumerate(lines):
                if i > 0:
                    p.add_run().add_break()
                p.add_run(line)
        else:
            doc.add_paragraph('Описание не заполнено.')

        doc.add_heading('Программа по дням', level=3)
        if e.program_entries:
            for entry in e.program_entries:
                header_parts = [entry.date_display or 'Дата не указана']
                if entry.time:
                    header_parts.append(entry.time)
                if entry.location:
                    header_parts.append(entry.location)
                header_p = doc.add_paragraph(style='List Bullet')
                header_p.add_run(' • '.join(header_parts)).bold = True
                if entry.description:
                    desc_p = doc.add_paragraph()
                    desc_p.paragraph_format.left_indent = Pt(18)
                    lines = entry.description.split('\n')
                    for i, line in enumerate(lines):
                        if i > 0:
                            desc_p.add_run().add_break()
                        desc_p.add_run(line)
        else:
            doc.add_paragraph('Нет данных.')

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'{safe_filename(title)}.docx', mimetype=DOCX_MIME)

def _festivals_report_xlsx(headers, rows, editions, title):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Фестивали'
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(str(h)) + 4)

    wrap = Alignment(wrap_text=True, vertical='top')

    ws2 = wb.create_sheet('Описания выпусков')
    ws2.append(['Выпуск', 'Описание'])
    for e in editions:
        ws2.append([f'{e.festival.name}, {e.year}', e.description or ''])
        ws2.cell(row=ws2.max_row, column=2).alignment = wrap
    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 70

    ws3 = wb.create_sheet('Программа по дням')
    ws3.append(['Выпуск', 'Программа'])
    for e in editions:
        label = f'{e.festival.name}, {e.year}'
        for entry in e.program_entries:
            header_parts = [entry.date_display or 'Дата не указана']
            if entry.time:
                header_parts.append(entry.time)
            if entry.location:
                header_parts.append(entry.location)
            cell_text = ' • '.join(header_parts)
            if entry.description:
                cell_text += '\n' + entry.description
            ws3.append([label, cell_text])
            ws3.cell(row=ws3.max_row, column=2).alignment = wrap
    ws3.column_dimensions['A'].width = 26
    ws3.column_dimensions['B'].width = 80

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'{safe_filename(title)}.xlsx', mimetype=XLSX_MIME)

# ---------- Фестивали: фильтруемый список выпусков ----------
@reports_bp.route('/festivals/export')
@login_required
def export_festivals():
    fmt = request.args.get('format', 'xlsx')
    name_q = request.args.get('name', '').strip()
    year_from = request.args.get('year_from', '').strip()
    year_to = request.args.get('year_to', '').strip()
    status = request.args.get('status', '').strip()

    editions = FestivalEdition.query.order_by(FestivalEdition.year.desc()).all()

    def matches(e):
        if name_q and name_q.lower() not in e.festival.name.lower():
            return False
        if year_from and e.year < int(year_from):
            return False
        if year_to and e.year > int(year_to):
            return False
        if status and e.status_name != status:
            return False
        return True

    filtered = [e for e in editions if matches(e)]

    def program_days_count(e):
        days = {(p.day, p.month) for p in e.program_entries if p.day and p.month}
        return len(days) if days else len(e.program_entries)

    headers = ['№', 'Название фестиваля', 'Год', 'Статус', 'Количество дней программы']
    rows = [[
        i + 1, e.festival.name, e.year, e.status_name, program_days_count(e),
    ] for i, e in enumerate(filtered)]

    title = 'Отчёт по фестивалям'
    if fmt == 'docx':
        return _festivals_report_docx(title, headers, rows, filtered)
    return _festivals_report_xlsx(headers, rows, filtered, title)
